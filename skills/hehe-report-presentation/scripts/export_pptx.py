"""Convert supported report HTML to native editable PowerPoint objects.

Text, rectangles, lines, tables and data-chart figures become native objects.
SVG illustrations remain individual images, listed in the conversion manifest.
"""
import argparse,io,json,re
from pathlib import Path
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE,PP_ALIGN,MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

# Pseudo-elements carry chapter panels and connector arrows in the report CSS.
# Materialize their computed appearance so the same native-object extractor sees them.
MATERIALIZE=r'''()=>{
const slide=document.querySelector('.slide.active');
if(slide.dataset.pptPrepared)return;
slide.dataset.pptPrepared='true';
if(!document.getElementById('ppt-pseudo-style')){
 const st=document.createElement('style');st.id='ppt-pseudo-style';
 st.textContent='[data-ppt-before]::before,[data-ppt-after]::after{content:none!important;display:none!important}';document.head.append(st);
}
for(const el of [slide,...slide.querySelectorAll('*')]){
 if(el.closest('svg,.native-chart'))continue;
 for(const side of ['before','after']){
  const css=getComputedStyle(el,'::'+side);
  if(css.content==='none'||css.content==='normal'||css.display==='none')continue;
  const text=css.content.replace(/^['"]|['"]$/g,'');
  if(!text.trim()&&css.backgroundColor==='rgba(0, 0, 0, 0)'&&parseFloat(css.borderTopWidth)===0)continue;
  const span=document.createElement('span');span.dataset.pptPseudo=side;
  for(const k of css)span.style.setProperty(k,css.getPropertyValue(k));
  span.textContent=text;el.setAttribute('data-ppt-'+side,'');
  side==='before'?el.prepend(span):el.append(span);
 }
}
}'''

EXTRACT=r'''()=>{
const s=document.querySelector('.slide.active'),root=s.getBoundingClientRect(),scale=root.width/s.offsetWidth;
const rect=e=>{const b=e.getBoundingClientRect();return {x:(b.x-root.x)/scale,y:(b.y-root.y)/scale,w:b.width/scale,h:b.height/scale}};
const style=e=>{const c=getComputedStyle(e);return {color:c.color,bg:c.backgroundColor,font:parseFloat(c.fontSize),bold:Number(c.fontWeight)>=600,align:c.textAlign,line:parseFloat(c.lineHeight)||parseFloat(c.fontSize)*1.4}};
const wrapped=nodes=>{const runs=[];let lineY=null;for(const node of nodes){const spec={...style(node.parentElement),href:node.parentElement.closest('a')?.href};for(let k=0;k<node.length;k++){let ch=node.textContent[k];const range=document.createRange();range.setStart(node,k);range.setEnd(node,k+1);const r=range.getBoundingClientRect();if(ch==='\n'||ch==='\r'){if(r.width>0&&['normal','nowrap'].includes(getComputedStyle(node.parentElement).whiteSpace))ch=' ';else continue;}if(!r.width&&!r.height)continue;const next=lineY!==null&&Math.abs(r.top-lineY)>spec.font*.7;if(next)runs.push({break:true});lineY=r.top;const last=runs.at(-1);if(last&&!last.break&&last.color===spec.color&&last.font===spec.font&&last.bold===spec.bold&&last.href===spec.href)last.text+=ch;else runs.push({...spec,text:ch});}}return runs;};
const textNodes=el=>{const nodes=[],walker=document.createTreeWalker(el,NodeFilter.SHOW_TEXT);let n;while(n=walker.nextNode())if(n.textContent.trim())nodes.push(n);return nodes;};
const visible=e=>e.getBoundingClientRect().height&&getComputedStyle(e).display!=='none';
const shapes=[];
for(const e of s.querySelectorAll('*')){
 if(!visible(e)||e.closest('svg,table,.native-chart'))continue;
 const c=getComputedStyle(e),r=rect(e);
 if(c.backgroundColor!=='rgba(0, 0, 0, 0)'&&c.backgroundColor!=='transparent')shapes.push({...r,fill:c.backgroundColor});
 if(e.tagName==='A')continue;
 for(const [side,x,y,w,h] of [['Top',r.x,r.y,r.w,0],['Bottom',r.x,r.y+r.h,r.w,0],['Left',r.x,r.y,0,r.h],['Right',r.x+r.w,r.y,0,r.h]]){
  if(parseFloat(c['border'+side+'Width'])>0&&c['border'+side+'Style']!=='none')shapes.push({x,y,w,h,line:c['border'+side+'Color'],weight:parseFloat(c['border'+side+'Width'])});
 }
}
const cellBar=c=>{
 const css=getComputedStyle(c),v=css.getPropertyValue('--val'),col=css.getPropertyValue('--score-color');
 if(v&&col)return {pct:parseFloat(v),color:col.trim(),tint:.17};
 const m=(c.getAttribute('style')||'').match(/linear-gradient\(to right,\s*(#[0-9a-f]+)\s+([\d.]+)%/i);
 return m?{pct:parseFloat(m[2]),color:m[1],tint:1}:null;
};
const tables=[...s.querySelectorAll('table')].map(t=>({...rect(t),rows:[...t.rows].map(row=>[...row.cells].map(c=>({...rect(c),...style(c),text:c.innerText,runs:wrapped(textNodes(c)),span:[c.colSpan,c.rowSpan],padding:parseFloat(getComputedStyle(c).paddingLeft),bar:cellBar(c)})))}));
const charts=[...s.querySelectorAll('.native-chart')].map(c=>({...rect(c),spec:JSON.parse(c.dataset.chart)}));
const blocks=new Map(),walker=document.createTreeWalker(s,NodeFilter.SHOW_TEXT);
let node;while(node=walker.nextNode()){
 if(!node.textContent.trim()||node.parentElement.closest('svg,table,.native-chart'))continue;
 let e=node.parentElement;
 while(e!==s&&['inline','contents'].includes(getComputedStyle(e).display))e=e.parentElement;
 if(!visible(e))continue;
 if(!blocks.has(e))blocks.set(e,[]);blocks.get(e).push(node);
}
const texts=[];
for(const [el,nodes] of blocks){
 const r=rect(el),runs=wrapped(nodes);
 const c=getComputedStyle(el),paddingLeft=parseFloat(c.paddingLeft),paddingTop=parseFloat(c.paddingTop);
 // Inline icons and flex labels have text bounds different from their parent.
 const range=document.createRange();range.setStart(nodes[0],0);range.setEnd(nodes.at(-1),nodes.at(-1).length);const b=range.getBoundingClientRect();
 const x=(b.x-root.x)/scale,y=(b.y-root.y)/scale;
 texts.push({x,y,w:Math.max(b.width/scale,r.x+r.w-x-parseFloat(c.paddingRight)),h:Math.max(b.height/scale,parseFloat(c.lineHeight)||22),...style(el),runs});
}
return {w:s.offsetWidth,h:s.offsetHeight,bg:getComputedStyle(s).backgroundColor,shapes,tables,charts,texts};}'''

def color(v):
    if v.startswith('#'):return RGBColor.from_string(v.lstrip('#'))
    nums=re.findall(r'[\d.]+',v)
    return RGBColor(*(int(float(x)) for x in nums[:3])) if len(nums)>=3 else RGBColor(30,55,70)
def pos(v):return Inches(v/96)
def font(run,spec):
    run.font.name='Microsoft YaHei';run.font.size=Pt(spec['font']*.75);run.font.bold=spec['bold'];run.font.color.rgb=color(spec['color'])
    if spec.get('href'):run.hyperlink.address=spec['href']
    props=run._r.get_or_add_rPr()
    ea=props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if ea is None:ea=OxmlElement('a:ea');props.append(ea)
    ea.set('typeface','Microsoft YaHei')

def paragraph_rules(para):
    props=para._p.get_or_add_pPr()
    props.set('eaLnBrk','1');props.set('latinLnBrk','0');props.set('hangingPunct','1')

def finish_links(prs,html_file):
    own=Path(html_file).resolve().as_uri()+'#page-'
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:continue
            runs=[r for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
            links=[r.hyperlink.address for r in runs]
            # A whole linked heading uses shape actions, preserving the source color.
            if links and links[0] and all(v==links[0] for v in links):
                link=links[0]
                for r in runs:r.hyperlink.address=None
                if link.startswith(own):
                    idx=int(link[len(own):])-1
                    if 0<=idx<len(prs.slides):sh.click_action.target_slide=prs.slides[idx]
                else:sh.click_action.hyperlink.address=link
    # Inline citations retain native text links in the report's muted blue.
    theme=next(r.target_part for r in prs.slide_masters[0].part.rels.values() if r.reltype==RT.THEME)
    root=etree.fromstring(theme.blob)
    for key in ['hlink','folHlink']:
        node=root.find('.//'+qn('a:'+key));node.clear()
        c=OxmlElement('a:srgbClr');c.set('val','34758B');node.append(c)
    theme._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)

def add_chart(slide,a):
    spec=a['spec'];kind=spec['chart_type'];data=CategoryChartData();data.categories=spec['labels']
    colors=['#27647f','#208e91','#b78647','#8ba6b2']
    if kind=='waterfall':
        values=spec['values'];totals=set(spec.get('totals',[0,len(values)-1]));current=0
        rows=[[],[],[],[]]
        for i,v in enumerate(values):
            if i in totals:
                rows[0].append(0);rows[1].append(0);rows[2].append(0);rows[3].append(v);current=v
            else:
                end=current+v;rows[0].append(min(current,end));rows[1].append(max(v,0));rows[2].append(max(-v,0));rows[3].append(0);current=end
        for name,vs in zip(['浮动基数','增加','减少','总额'],rows):data.add_series(name,vs)
        colors=['#ffffff','#238d88','#a6674d','#173f54']
    else:
        for ser in spec['series']:data.add_series(ser['name'],ser['values'])
    typ={'bar':XL_CHART_TYPE.COLUMN_CLUSTERED,'hbar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS,'stacked':XL_CHART_TYPE.COLUMN_STACKED,'waterfall':XL_CHART_TYPE.COLUMN_STACKED}[kind]
    chart=slide.shapes.add_chart(typ,pos(a['x']),pos(a['y']),pos(a['w']),pos(a['h']),data).chart
    chart.has_legend=False;chart.chart_style=2;chart.has_title=False
    chart.category_axis.tick_labels.font.name='Microsoft YaHei';chart.category_axis.tick_labels.font.size=Pt(10)
    chart.value_axis.tick_labels.font.size=Pt(10)
    if kind!='hbar' or min(v for s in spec['series'] for v in s['values'])>=0:chart.value_axis.minimum_scale=0
    chart.value_axis.major_gridlines.format.line.color.rgb=color('#d8e2e7');chart.value_axis.major_gridlines.format.line.width=Pt(.5)
    chart.category_axis.format.line.color.rgb=color('#cfdee4');chart.value_axis.format.line.fill.background()
    if kind=='hbar':chart.category_axis.reverse_order=True
    if kind in {'bar','hbar'}:
        plot=chart.plots[0];plot.has_data_labels=True;plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.size=Pt(11)
    for j,ser in enumerate(chart.series):
        fill=ser.format.fill;fill.solid();fill.fore_color.rgb=color(colors[j%4]);ser.format.line.fill.background()
        if kind=='waterfall':
            if j==0:ser.format.fill.background()
            else:
                # Label only visible contributions, using the supported center position.
                for i,point in enumerate(ser.points):
                    if not rows[j][i]:continue
                    label=point.data_label;label.position=XL_LABEL_POSITION.CENTER
                    label.text_frame.text=f'{values[i]:g}' if i in totals else f'{values[i]:+g}'
                    if i not in totals:
                        layout=OxmlElement('c:layout');manual=OxmlElement('c:manualLayout')
                        for tag,val in [('xMode','factor'),('yMode','factor'),('x','0'),('y','-0.06')]:
                            item=OxmlElement('c:'+tag);item.set('val',val);manual.append(item)
                        layout.append(manual);label._get_or_add_dLbl().insert(1,layout)
                    for para in label.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.name='Microsoft YaHei';run.font.size=Pt(10);run.font.color.rgb=color('#ffffff' if i in totals else '#173f54')
    if spec.get('max'):chart.value_axis.maximum_scale=spec['max']
    return chart

def convert(html_file,output,browser_path,replace=False):
    output=Path(output).resolve();output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists() and not replace:raise FileExistsError(output)
    prs=Presentation();manifest={'source':str(Path(html_file).resolve()),'pages':[],'image_elements':[],'omitted_decorations':[],'text_wrapping':'browser line breaks retained in native paragraphs'}
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path=browser_path,headless=True)
        page=browser.new_page(viewport={'width':1600,'height':1500});page.context.set_offline(True)
        page.goto(Path(html_file).resolve().as_uri());page.wait_for_function("document.documentElement.dataset.ready==='true'")
        if page.locator('body').get_attribute('data-canvas')=='scroll':raise ValueError('PPT需要分页源；先按主题编排下拉文稿')
        for n in range(page.locator('.slide').count()):
            page.evaluate('(n)=>show(n)',n);page.evaluate(MATERIALIZE);d=page.evaluate(EXTRACT)
            prs.slide_width=pos(d['w']);prs.slide_height=pos(d['h']);slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(d['bg'])
            for a in d['shapes']:
                if a['w']<=0 and a['h']<=0:continue
                shp=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,pos(max(0,a['x'])),pos(max(0,a['y'])),pos(max(.5,a['w'])),pos(max(.5,a['h'])))
                shp.line.fill.background();shp.fill.solid();shp.fill.fore_color.rgb=color(a.get('fill',a.get('line')))
                shp._element.spPr.append(OxmlElement('a:effectLst'))
            for a in d['tables']:
                if any(c['span']!=[1,1] for row in a['rows'] for c in row):raise ValueError('合并单元格需在页面源中展开后转换')
                for row in a['rows']:
                    for c in row:
                        if not c['bar']:continue
                        b=c['bar'];base=color(b['color']);t=b['tint']
                        tint=RGBColor(*(round(v*t+255*(1-t)) for v in base))
                        for fraction,fill in [(1,color('#ffffff')),(max(0,min(100,b['pct']))/100,tint)]:
                            rect=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,pos(c['x']),pos(c['y']),pos(c['w']*fraction),pos(c['h']))
                            rect.fill.solid();rect.fill.fore_color.rgb=fill;rect.line.fill.background()
                            rect._element.spPr.append(OxmlElement('a:effectLst'))
                            rect.name='表格数值背景条' if fraction!=1 else '表格背景'
                table=slide.shapes.add_table(len(a['rows']),len(a['rows'][0]),pos(a['x']),pos(a['y']),pos(a['w']),pos(a['h'])).table
                for j,c in enumerate(a['rows'][0]):table.columns[j].width=pos(c['w'])
                for i,row in enumerate(a['rows']):
                    table.rows[i].height=pos(max(c['h'] for c in row))
                    for j,c in enumerate(row):
                        cell=table.cell(i,j);cell.margin_left=cell.margin_right=pos(c['padding']);cell.margin_top=pos(7);cell.margin_bottom=pos(3);cell.vertical_anchor=MSO_ANCHOR.TOP
                        cell.fill.solid();cell.fill.fore_color.rgb=color(c['bg'] if c['bg']!='rgba(0, 0, 0, 0)' else ('#edf4f6' if i%2==0 else '#ffffff'))
                        if c['bar']:cell.fill.background()
                        tf=cell.text_frame;tf.clear();tf.word_wrap=False
                        para=tf.paragraphs[0]
                        for spec in c['runs']:
                            if spec.get('break'):para.add_line_break()
                            else:
                                run=para.add_run();run.text=spec['text'];font(run,spec)
                        for para in tf.paragraphs:
                            paragraph_rules(para)
                            para.space_before=Pt(0);para.space_after=Pt(0);para.line_spacing=Pt(c['line']*.75)
                            para.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(c['align'],PP_ALIGN.LEFT)
            for a in d['charts']:
                add_chart(slide,a)
            for a in d['texts']:
                if a['w']<1 or a['h']<1:continue
                shape=slide.shapes.add_textbox(pos(a['x']),pos(a['y']-1),pos(a['w']+2),pos(a['h']+5));tf=shape.text_frame
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
                para=tf.paragraphs[0];para.space_before=Pt(0);para.space_after=Pt(0);para.line_spacing=Pt(a['line']*.75)
                paragraph_rules(para)
                para.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(a['align'],PP_ALIGN.LEFT)
                for spec in a['runs']:
                    if spec.get('break'):para.add_line_break()
                    else:
                        run=para.add_run();run.text=spec['text'];font(run,spec)
            svgs=page.locator('.slide.active svg:not(.native-chart svg)')
            for i in range(svgs.count()):
                el=svgs.nth(i)
                if el.evaluate("e=>!!e.closest('.native-chart')"):continue
                b=el.bounding_box();root=page.locator('.slide.active').bounding_box();sc=root['width']/d['w']
                if b['width']<=0 or b['height']<=0:continue
                x=max(b['x'],root['x']);y=max(b['y'],root['y'])
                w=min(b['x']+b['width'],root['x']+root['width'])-x;h=min(b['y']+b['height'],root['y']+root['height'])-y
                if w<=0 or h<=0:continue
                # Keep the visible part of intentionally cropped cover artwork.
                blob=page.screenshot(clip={'x':x,'y':y,'width':w,'height':h})
                picture=slide.shapes.add_picture(io.BytesIO(blob),pos((x-root['x'])/sc),pos((y-root['y'])/sc),pos(w/sc),pos(h/sc))
                # Artwork may share space with a title; keep editable text above it.
                tree=slide.shapes._spTree;tree.remove(picture._element)
                first_text=next((i for i,e in enumerate(tree) if e.tag.endswith('}sp') and e.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}t') is not None),len(tree)-1)
                tree.insert(first_text,picture._element)
                manifest['image_elements'].append({'page':n+1,'element':'SVG插图','editable':'独立图片，可替换；图内路径不可逐点编辑'})
            manifest['pages'].append({'page':n+1,'text_boxes':len(d['texts']),'tables':len(d['tables']),'charts':len(d['charts']),'shapes':sum(x.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE for x in slide.shapes)})
        browser.close()
    finish_links(prs,html_file)
    prs.save(output)
    manifest['reopened_pages']=len(Presentation(output).slides)
    output.with_name(output.stem+'_对象核验_CO.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pages':len(prs.slides),'tables':sum(x['tables'] for x in manifest['pages']),'charts':sum(x['charts'] for x in manifest['pages']),'text_boxes':sum(x['text_boxes'] for x in manifest['pages']),'image_elements':len(manifest['image_elements'])},ensure_ascii=False))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('html');ap.add_argument('output');ap.add_argument('--browser',required=True);ap.add_argument('--replace',action='store_true');a=ap.parse_args();convert(a.html,a.output,a.browser,a.replace)
